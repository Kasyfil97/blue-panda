"""
hierarchy_v3_eval.py -- evaluates the manager's REVISED resolver hierarchy:
a flat, resolver-level cascade (not grouped into layers), reordered based on
each resolver's own ISOLATED accuracy from the v2 evaluation:

    1. exact_match       ExactMatchResolver (unchanged)
    2. kata_similar       KataSimilarResolver (BM25 context still applied
                           conditionally for AS400 tables, same as v2)
    3. kata_alias          KataEvidenceResolver
    4. bm25_informatica    BM25Resolver (informatica_certified priority)
    5. bm25_confluence     BM25Resolver (confluence source_type, BM25 index --
                           NOT live Confluence page search)
    6. bm25_as400          BM25Resolver (AS400 tables only)
    7. llm                 PureLLMResolver (last resort)

Why reordered: in v2, kata_alias ran BEFORE kata_similar inside the "KATA"
layer (first-success-wins), even though kata_alias's own isolated accuracy
(81-83% similar) is meaningfully lower than kata_similar's apparent accuracy
(89.6-89.7% similar in the v2 Resolver Breakdown). That meant kata_alias was
"stealing" rows kata_similar might have answered better, capping the
combined KATA layer's blended accuracy at ~85% instead of closer to
kata_similar's own number. Swapping the order lets the stronger resolver
claim rows first.

IMPORTANT CAVEAT ON THAT v2 "kata_similar" NUMBER: in v2, kata_similar was
only ever TRIED (and only ever measured) on the subset of rows where
kata_alias had ALREADY failed (~159 of 196 rows) -- never on the full 196 in
isolation, because it only ran in the "else" branch after kata_alias's
attempt. Its true isolated accuracy on ALL 196 rows (including the 37 rows
kata_alias used to grab first) is measured for the FIRST TIME in this
script, since kata_similar now runs unconditionally, first, on every row.
Read the new Ranking sheet's kata_similar row as the real number; the v2
figure was a lower-bound estimate on a biased subset.

Confluence (the free-text page-search resolvers: confluence_similar and the
regex-based confluence_fallback) is dropped entirely from this hierarchy --
in the v2 full run both underperformed plain PureLLM (34-38% similar vs
PureLLM's 39.5%) -- replaced here by bm25_confluence (BM25-indexed
Confluence-sourced content, a different and already-curated source from the
live Confluence page search) and llm as the final fallback.

Method: every one of the 7 resolvers runs ISOLATED on every applicable row
(same principle as hierarchy_v2's layer_eval.py), so each step's own
coverage/accuracy is measured independently; the sequential cascade
(Waterfall sheet) is then DERIVED from those isolated results -- identical
to what a real first-hit-wins run would produce, without re-running
anything.

Usage:
    python hierarchy_v3_eval.py --sample 20   # smoke test
    python hierarchy_v3_eval.py               # full run
    python hierarchy_v3_eval.py --resume-raw evaluation/hierarchy_v3_raw_XXXX.json
"""

from __future__ import annotations

import sys
from pathlib import Path


def _find_project_root(start: Path) -> Path:
    for candidate in [start, *start.parents]:
        if (candidate / "mage_flow").is_dir():
            return candidate
    raise RuntimeError(
        f"Could not find a 'mage_flow' package above {start} -- "
        "this script must live somewhere inside (or under) the blue-panda project root."
    )


_SCRIPT_DIR = Path(__file__).resolve().parent
_PROJECT_ROOT = _find_project_root(_SCRIPT_DIR)
for _p in (str(_PROJECT_ROOT), str(_SCRIPT_DIR)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import argparse
import asyncio
import datetime as dt
import json
import logging
import os
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from tqdm import tqdm

import evaluate_column_description as ecd
import layer_metrics as lm
from mage_flow import config
from mage_flow.clients import BM25Client, get_llm_client
from mage_flow.common import is_missing_desc, norm
from mage_flow.kata_similar import KataDatasetPool, KataSimilarResolver, extract_keyword, retrieve_bm25_knowledge
from mage_flow.llm_generation import LLMGeneration
from mage_flow.resolvers import (
    _AS400_SOURCE_TYPES,
    BM25Resolver,
    ExactMatchResolver,
    KataEvidenceResolver,
    PureLLMResolver,
    ResolverContext,
    _is_as400_table,
)
from mage_flow.source_priority import KNOWLEDGE_PRIORITY_AS400, KNOWLEDGE_PRIORITY_INFORMATICA_CERTIFIED

BASE_DIR = _PROJECT_ROOT
TABLE_COL = "Tabel"
COLUMN_COL = "Kolom"
DATA_TYPE_COL = "Tipe Data"
GT_TITLE_COL = "gt business title"
GT_DESC_COL = "gt column description"
PRED_DESC_COL = "Predicted Column Description"
PRED_TITLE_COL = "Predicted Business Title"

DEFAULT_INPUT_CANDIDATES = [
    "test set eval v2 mage.xlsx", "test set eval v2.xlsx", "test set eval 6.xlsx", "test_set_eval_v2.xlsx",
]
JUDGE_DEPLOYMENT = os.getenv("AZURE_OPENAI_DEPLOYMENT_FULL", "gpt-4.1")

# The new flat, resolver-level hierarchy, in order. This IS lm.LAYER_ORDER
# for this script -- each "layer" is exactly one resolver now, so
# layer_metrics.py's existing summary/ranking/waterfall logic applies
# unchanged; no per-layer resolver LIST / internal first-success-wins
# sub-cascade is needed anymore (that nesting was v2's design).
STEP_ORDER: List[str] = [
    "exact_match", "kata_similar", "kata_alias",
    "bm25_informatica", "bm25_confluence", "bm25_as400", "llm",
]
STEP_LABELS: Dict[str, str] = {
    "exact_match": "1. Exact Match (BM25 exact lookup)",
    "kata_similar": "2. KATA Similar (LLM selection; BM25 context on AS400 tables)",
    "kata_alias": "3. KATA Alias (exact alias / technical relation)",
    "bm25_informatica": "4. BM25 Informatica (certified priority)",
    "bm25_confluence": "5. BM25 Confluence (BM25-indexed, not live page search)",
    "bm25_as400": "6. BM25 AS400 (AS400 tables only)",
    "llm": "7. Pure LLM (last resort)",
}
TRUTHFULNESS_STEPS = ("kata_similar",)  # same scope as v2 -- only KATA's LLM selection is audited

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
logger = logging.getLogger("hierarchy_v3_eval")


# ---------------------------------------------------------------------------
# Resolver construction
# ---------------------------------------------------------------------------

def build_resolvers(pool: KataDatasetPool) -> Dict[str, Any]:
    return {
        "exact_match": ExactMatchResolver(allowed_priorities={KNOWLEDGE_PRIORITY_AS400}, resolution_tag="exact_as400"),
        "kata_similar": KataSimilarResolver(pool=pool, use_bm25_context=True),
        "kata_alias": KataEvidenceResolver(),
        "bm25_informatica": BM25Resolver(
            allowed_priorities={KNOWLEDGE_PRIORITY_INFORMATICA_CERTIFIED},
            source_types=["informatica"], resolution_tag="bm25_informatica_certified",
        ),
        "bm25_confluence": BM25Resolver(allowed_priorities=None, source_types=["confluence"], resolution_tag="bm25_confluence"),
        "bm25_as400": BM25Resolver(
            allowed_priorities=None, source_types=_AS400_SOURCE_TYPES,
            run_if=lambda ctx: _is_as400_table(ctx.table_name), resolution_tag="bm25_as400",
        ),
        "llm": PureLLMResolver(),
    }


def run_one(resolver: Any, ctx: ResolverContext) -> Tuple[Optional[Any], str]:
    """Run a single resolver; '' description counts as not-covered (mirrors
    PureLLMResolver's 'unknown' case and v2's run_resolvers())."""
    try:
        result = resolver.resolve(ctx)
    except Exception as exc:  # noqa: BLE001
        logger.warning("%s failed on %s.%s: %s", type(resolver).__name__, ctx.table_name, ctx.col_name, exc)
        return None, f"{type(resolver).__name__}: {exc}"
    if result.resolved and not is_missing_desc(result.description or ""):
        return result, ""
    return None, ""


def summarize_knowledge(result: Optional[Any], max_items: int = 5, max_chars: int = 300) -> str:
    """Human-readable 'what evidence fed this prediction' string for the
    Step Detail sheet -- generic across all resolver types, since every
    ResolverResult.knowledge item carries at least column_description/
    source_type (the BaseResolver contract used throughout this codebase).

    kata_similar is a special case: the AS400 BM25 evidence that actually
    informed its candidate SELECTION lives nested under
    knowledge[0]["kata_similar"]["bm25_evidence"] (see kata_similar.py),
    separate from the selected candidate's own column_description -- both
    are surfaced here so "what evidence fed this decision" is complete.
    """
    if not result or not result.knowledge:
        return ""
    parts = []
    for item in result.knowledge[:max_items]:
        desc = norm(item.get("column_description"))
        src = norm(item.get("source_type"))
        name = norm(item.get("data_element_name") or item.get("business_title"))
        label = f"[{src}]" if src else ""
        bits = " ".join(b for b in (label, name, desc) if b)
        if bits:
            parts.append(bits[:max_chars])
        bm25_evidence = norm((item.get("kata_similar") or {}).get("bm25_evidence"))
        if bm25_evidence:
            parts.append(f"[bm25_context] {bm25_evidence}"[:max_chars])
    return " || ".join(parts)


def derive_title(result: Any, ctx: ResolverContext, llm: LLMGeneration, bt_sampling: Optional[dict]) -> str:
    for item in result.knowledge or []:
        title = norm(item.get("business_title"))
        if title:
            return title
    for item in result.knowledge or []:
        if norm(item.get("source_type")).lower() == "kata" and norm(item.get("data_element_name")):
            return norm(item.get("data_element_name"))
    try:
        from mage_flow.flow import _generate_business_title
        return _generate_business_title(
            table_name=ctx.table_name, col_name=ctx.col_name, col_description=norm(result.description),
            system_context=ctx.system_context, knowledge=result.knowledge, llm=llm, sampling_params=bt_sampling,
        )
    except Exception as exc:  # noqa: BLE001
        logger.warning("business title generation failed for %s.%s: %s", ctx.table_name, ctx.col_name, exc)
        return ""


def to_outcome(result: Optional[Any], error: str, ctx: ResolverContext, llm: LLMGeneration, bt_sampling: Optional[dict]) -> Dict[str, Any]:
    if result is None:
        return {"covered": False, "tag": "", "desc": "", "title": "", "reason": "", "error": error, "knowledge": ""}
    first = (result.knowledge or [{}])[0]
    return {
        "covered": True,
        "tag": result.resolution_tag,
        "desc": norm(result.description),
        "title": derive_title(result, ctx, llm, bt_sampling),
        "reason": norm(first.get("match_reason")),
        "error": "",
        "knowledge": summarize_knowledge(result),
    }


# ---------------------------------------------------------------------------
# Per-table prep -- no Confluence page-search prep needed (dropped entirely)
# ---------------------------------------------------------------------------

class Deps:
    def __init__(self, bm25: BM25Client, llm: LLMGeneration, settings: Dict[str, Any], pool: KataDatasetPool) -> None:
        self.bm25 = bm25
        self.llm = llm
        self.settings = settings
        self.pool = pool
        self.sampling = settings.get("llm", {}).get("sampling", {"temperature": 0.1, "top_p": 0.9, "max_tokens": 1200})
        bm25_cfg = settings.get("bm25", {})
        table_cfg = bm25_cfg.get("table_search", {})
        self.bm25_params = {
            "table": {"top_k": int(table_cfg.get("top_k", 20)), "threshold": float(table_cfg.get("threshold", 8.0))},
            "term": bm25_cfg.get("term_search", {}),
            "global": bm25_cfg.get("global_search", {}),
        }
        self.bt_sampling = settings.get("business_title", {}).get("sampling")


def prepare_table(table: str, columns: List[str], deps: Deps) -> Dict[str, Any]:
    request_id = str(uuid.uuid4())
    payload = deps.bm25.get_system_context(table_name=table, request_id=request_id)
    payload = payload if isinstance(payload, dict) else {}
    table_term = deps.bm25.get_table_term_context(table, request_id=request_id) or []
    table_hits = deps.bm25.get_table_knowledge(
        table_name=table, top_k=deps.bm25_params["table"]["top_k"],
        threshold=deps.bm25_params["table"]["threshold"], request_id=request_id,
    ) or []
    # Warm the KATA dataset cache up front (one search per table, shared
    # across its columns -- same principle as v2).
    deps.pool.get(table, extract_keyword_fn=lambda raw: extract_keyword(deps.llm.client, raw))
    logger.info("prepared %s: as400=%s table_terms=%d", table, _is_as400_table(table), len(table_term))
    return {
        "request_id": request_id, "system_context": payload.get("context", "") or "",
        "table_term": table_term, "table_hits": table_hits,
    }


# ---------------------------------------------------------------------------
# Per-row execution -- all 7 resolvers run ISOLATED on every row
# ---------------------------------------------------------------------------

def process_row(ri: int, row: pd.Series, prep: Dict[str, Any], resolvers: Dict[str, Any], deps: Deps) -> Dict[str, Any]:
    table, col = norm(row[TABLE_COL]), norm(row[COLUMN_COL])
    tp = prep[table]
    ctx = ResolverContext(
        table_name=table, col_name=col, system_context=tp["system_context"], table_hits=tp["table_hits"],
        bm25=deps.bm25, llm=deps.llm, sampling_params=deps.sampling, bm25_params=deps.bm25_params,
        settings=deps.settings, request_id=tp["request_id"], table_term_context=tp["table_term"],
    )
    ctx.data_type = norm(row.get(DATA_TYPE_COL))

    # AS400 BM25 evidence -- feeds kata_similar's context (same as v2) AND
    # is what bm25_as400's own isolated attempt below also draws on.
    try:
        evidence = retrieve_bm25_knowledge(resolvers["bm25_as400"], ctx)
    except Exception as exc:  # noqa: BLE001
        logger.warning("BM25 AS400 evidence retrieval failed for %s.%s: %s", table, col, exc)
        evidence = []
    ctx.extra_knowledge = evidence

    outcomes: Dict[str, Dict[str, Any]] = {}
    for step in STEP_ORDER:
        result, error = run_one(resolvers[step], ctx)
        outcomes[step] = to_outcome(result, error, ctx, deps.llm, deps.bt_sampling)

    return {"row_id": ri, "is_as400": _is_as400_table(table), "n_evidence": len(evidence), "outcomes": outcomes}


def run_all_rows(df: pd.DataFrame, deps: Deps, workers: int) -> List[Dict[str, Any]]:
    resolvers = build_resolvers(deps.pool)
    columns_by_table: Dict[str, List[str]] = {}
    for _, row in df.iterrows():
        columns_by_table.setdefault(norm(row[TABLE_COL]), []).append(norm(row[COLUMN_COL]))

    prep: Dict[str, Any] = {}
    for table, cols in tqdm(columns_by_table.items(), desc="Preparing tables"):
        prep[table] = prepare_table(table, cols, deps)

    results: List[Optional[Dict[str, Any]]] = [None] * len(df)
    with ThreadPoolExecutor(max_workers=max(1, workers)) as executor:
        futures = {executor.submit(process_row, ri, row, prep, resolvers, deps): ri for ri, row in df.iterrows()}
        for future in tqdm(as_completed(futures), total=len(futures), desc="Running hierarchy v3"):
            ri = futures[future]
            try:
                results[ri] = future.result()
            except Exception as exc:  # noqa: BLE001
                logger.error("row %d failed entirely: %s", ri, exc)
                empty = {"covered": False, "tag": "", "desc": "", "title": "", "reason": "", "error": str(exc), "knowledge": ""}
                results[ri] = {"row_id": ri, "is_as400": _is_as400_table(norm(df.at[ri, TABLE_COL])),
                               "n_evidence": 0, "outcomes": {s: dict(empty) for s in STEP_ORDER}}
    return [r for r in results if r is not None]


# ---------------------------------------------------------------------------
# Long table, judging (same pattern/machinery as v2's layer_eval.py)
# ---------------------------------------------------------------------------

def build_long_df(df: pd.DataFrame, results: List[Dict[str, Any]]) -> pd.DataFrame:
    records = []
    for res in results:
        ri = res["row_id"]
        for step in STEP_ORDER:
            o = res["outcomes"][step]
            records.append({
                "row_id": ri, "Tabel": norm(df.at[ri, TABLE_COL]), "Kolom": norm(df.at[ri, COLUMN_COL]),
                "is_as400": res["is_as400"], "layer": step, "covered": bool(o["covered"]), "tag": o["tag"],
                "pred_desc": o["desc"], "pred_title": o["title"], "reason": o["reason"], "error": o["error"],
                "knowledge": o["knowledge"], "n_candidates": None,
                "gt_desc": df.at[ri, GT_DESC_COL], "gt_title": df.at[ri, GT_TITLE_COL],
                "eval_desc": "", "eval_title": "", "truthfulness": "", "truthfulness_reason": "",
            })
    return pd.DataFrame(records)


def judge_similarity(long_df: pd.DataFrame) -> pd.DataFrame:
    covered_idx = long_df.index[long_df["covered"]]
    if len(covered_idx) == 0:
        return long_df

    def labels_for(gt_col, pred_col, gt_name, pred_name, field, singular) -> Dict[int, str]:
        out: Dict[int, str] = {}
        sendable = []
        for i in covered_idx:
            gt_missing = is_missing_desc(norm(long_df.at[i, gt_col]))
            pred_missing = is_missing_desc(norm(long_df.at[i, pred_col]))
            if gt_missing:
                out[i] = "no_gt"
            elif pred_missing:
                out[i] = "no_pred"
            else:
                sendable.append(i)
        if sendable:
            frame = pd.DataFrame({
                gt_name: [long_df.at[i, gt_col] for i in sendable],
                pred_name: [long_df.at[i, pred_col] for i in sendable],
            })
            labels = asyncio.run(ecd.evaluate_all(frame, gt_name, pred_name, field, singular, deployment=JUDGE_DEPLOYMENT))
            out.update(dict(zip(sendable, labels)))
        return out

    desc = labels_for("gt_desc", "pred_desc", GT_DESC_COL, PRED_DESC_COL, "column descriptions", "column description")
    title = labels_for("gt_title", "pred_title", GT_TITLE_COL, PRED_TITLE_COL, "business titles", "business title")
    long_df.loc[list(desc), "eval_desc"] = pd.Series(desc)
    long_df.loc[list(title), "eval_title"] = pd.Series(title)
    return long_df


TRUTHFULNESS_SYSTEM_PROMPT = """\
You are an independent auditor checking whether a metadata-matching decision
was genuinely justified, or fabricated a connection that isn't really there.

You did NOT make this decision -- a different system selected a KATA data
element as the match for a database column. Your job is only to verify it.

INPUTS:
- Table name: {table_name}
- Column name: {col_name}
- Column data type: {data_type}
- Selected candidate / retrieved evidence: {candidate_info}

TASK:
Judge whether the selection is genuinely, directly supported by the
information given -- not a stretch, not a superficial wording coincidence,
and not reliant on an assumption about what an abbreviation "probably"
means when that assumption isn't stated anywhere in the inputs above.

OUTPUT FORMAT (STRICT JSON ONLY):
{{"truthful": true, "reasoning": "<one short sentence, in Indonesian>"}}
"""


async def truthfulness_batch(items: List[Dict[str, str]]) -> List[Tuple[str, str]]:
    from openai import AsyncAzureOpenAI
    client = AsyncAzureOpenAI(azure_endpoint=ecd.AZURE_ENDPOINT, api_key=ecd.AZURE_API_KEY, api_version=ecd.AZURE_API_VER)
    semaphore = asyncio.Semaphore(ecd.CONCURRENCY)

    async def one(item: Dict[str, str]) -> Tuple[str, str]:
        prompt = TRUTHFULNESS_SYSTEM_PROMPT.format(**item)
        messages = [{"role": "system", "content": prompt}, {"role": "user", "content": "Assess now."}]
        async with semaphore:
            for attempt in range(1, 4):
                try:
                    resp = await client.chat.completions.create(model=JUDGE_DEPLOYMENT, messages=messages, temperature=0, max_tokens=250)
                    raw = resp.choices[0].message.content.strip().replace("```json", "").replace("```", "").strip()
                    parsed = json.loads(raw)
                    return ("true" if bool(parsed.get("truthful")) else "false"), norm(parsed.get("reasoning"))
                except Exception as exc:  # noqa: BLE001
                    if attempt == 3:
                        return "error", str(exc)
                    await asyncio.sleep(2 * attempt)
        return "error", "unreachable"

    return list(await asyncio.gather(*[one(item) for item in items]))


def judge_truthfulness(long_df: pd.DataFrame, df: pd.DataFrame) -> pd.DataFrame:
    idx = long_df.index[long_df["covered"] & long_df["layer"].isin(TRUTHFULNESS_STEPS)]
    if len(idx) == 0:
        return long_df
    items = []
    for i in idx:
        row = long_df.loc[i]
        items.append({
            "table_name": row["Tabel"], "col_name": row["Kolom"],
            "data_type": norm(df.at[row["row_id"], DATA_TYPE_COL]) or "(unknown)",
            "candidate_info": row["knowledge"] or "(none)",
        })
    verdicts = asyncio.run(truthfulness_batch(items))
    long_df.loc[idx, "truthfulness"] = [v[0] for v in verdicts]
    long_df.loc[idx, "truthfulness_reason"] = [v[1] for v in verdicts]
    return long_df


# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------

NOTES = [
    ("What changed vs v2", "The hierarchy is now a FLAT, resolver-level cascade (not resolvers grouped inside 4 layers). kata_similar now runs BEFORE kata_alias (v2 had it the other way), and Confluence page-search resolvers are dropped entirely -- replaced by bm25_confluence + llm."),
    ("Why kata_similar moved up", "In v2, kata_alias ran first and had lower isolated accuracy (81-83% similar) than kata_similar (89.6-89.7%, but only measured on the ~159 rows kata_alias had already failed). Trying the stronger resolver first should raise the blended accuracy."),
    ("kata_similar's number here is NEW information", "This is the first time kata_similar has been measured on ALL rows unconditionally (v2 only tested it on kata_alias's failures). Compare this Ranking row to v2's Resolver Breakdown 'kata_similar' row -- if they differ, that's a real methodological finding, not noise."),
    ("Method", "Every one of the 7 resolvers runs ISOLATED on every row (same principle as v2), so each step's coverage/accuracy is measured independently. Waterfall is DERIVED from those isolated results in the new order -- equivalent to a real first-hit-wins run."),
    ("Step Detail sheet", "One row per (test row, step): the knowledge/evidence retrieved, ground truth, prediction, and verdict -- for auditing exactly what each step saw and produced."),
    ("Truthfulness", "Independent GPT-4.1 audit (True/False), kata_similar only -- same scope and strict standard as v2."),
    ("Judge / generator", f"Generator: LLM_PROVIDER endpoint. Judge: {JUDGE_DEPLOYMENT}."),
]


def wide_results(df: pd.DataFrame, long_df: pd.DataFrame) -> pd.DataFrame:
    winners = lm.winner_by_row(long_df, STEP_ORDER)
    base = df[[TABLE_COL, COLUMN_COL, DATA_TYPE_COL, GT_TITLE_COL, GT_DESC_COL]].copy() if DATA_TYPE_COL in df.columns \
        else df[[TABLE_COL, COLUMN_COL, GT_TITLE_COL, GT_DESC_COL]].copy()
    base["AS400 table"] = [_is_as400_table(norm(t)) for t in base[TABLE_COL]]
    base["Cascade winner"] = [winners.get(i, "") for i in base.index]
    for step in STEP_ORDER:
        sub = long_df[long_df["layer"] == step].set_index("row_id")
        prefix = f"[{step}]"
        base[f"{prefix} covered"] = [bool(sub["covered"].get(i, False)) for i in base.index]
        base[f"{prefix} resolver"] = [sub["tag"].get(i, "") for i in base.index]
        base[f"{prefix} description"] = [sub["pred_desc"].get(i, "") for i in base.index]
        base[f"{prefix} business title"] = [sub["pred_title"].get(i, "") for i in base.index]
        base[f"{prefix} eval desc"] = [sub["eval_desc"].get(i, "") for i in base.index]
        base[f"{prefix} eval title"] = [sub["eval_title"].get(i, "") for i in base.index]
        if step in TRUTHFULNESS_STEPS:
            base[f"{prefix} truthfulness"] = [sub["truthfulness"].get(i, "") for i in base.index]
            base[f"{prefix} truthfulness reasoning"] = [sub["truthfulness_reason"].get(i, "") for i in base.index]
    return base


def step_detail(df: pd.DataFrame, long_df: pd.DataFrame) -> pd.DataFrame:
    """Long-format: one row per (test row, step) -- knowledge retrieved, GT,
    prediction, verdict, side by side. This is the granular view requested."""
    rows = []
    for step in STEP_ORDER:
        sub = long_df[long_df["layer"] == step]
        for _, r in sub.iterrows():
            rows.append({
                "Step": STEP_LABELS.get(step, step), "Tabel": r["Tabel"], "Kolom": r["Kolom"],
                "Covered": r["covered"], "Resolver Tag": r["tag"],
                "Knowledge Retrieved": r["knowledge"],
                "GT Description": r["gt_desc"], "Predicted Description": r["pred_desc"], "Eval Description": r["eval_desc"],
                "GT Business Title": r["gt_title"], "Predicted Business Title": r["pred_title"], "Eval Business Title": r["eval_title"],
                "Truthfulness": r["truthfulness"], "Truthfulness Reasoning": r["truthfulness_reason"],
            })
    out = pd.DataFrame(rows)
    step_rank = {STEP_LABELS.get(s, s): i for i, s in enumerate(STEP_ORDER)}
    out["_order"] = out["Step"].map(step_rank)
    out = out.sort_values(["_order", "Tabel", "Kolom"]).drop(columns=["_order"]).reset_index(drop=True)
    return out


def write_workbook(path: Path, sheets: Dict[str, pd.DataFrame]) -> Path:
    from openpyxl.styles import Alignment, Font, PatternFill

    def _write(target: Path) -> None:
        with pd.ExcelWriter(target, engine="openpyxl") as writer:
            for name, frame in sheets.items():
                frame.to_excel(writer, sheet_name=name, index=False)
                ws = writer.sheets[name]
                ws.freeze_panes = "A2"
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                for idx, col in enumerate(frame.columns, start=1):
                    letter = ws.cell(row=1, column=idx).column_letter
                    is_pct = "%" in str(col) or "Yield" in str(col)
                    ws.column_dimensions[letter].width = 16 if is_pct else min(50, max(12, len(str(col)) + 2))
                    if is_pct:
                        for r in range(2, len(frame) + 2):
                            ws.cell(row=r, column=idx).number_format = "0.0%"
            if "Notes" in sheets:
                writer.sheets["Notes"].column_dimensions["A"].width = 30
                writer.sheets["Notes"].column_dimensions["B"].width = 120
                for row in writer.sheets["Notes"].iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")

    try:
        _write(path)
        return path
    except PermissionError:
        fallback = path.with_name(f"{path.stem}_{dt.datetime.now():%Y%m%d_%H%M%S}{path.suffix}")
        logger.warning("Could not write %s (open in Excel?) -- saving to %s instead.", path, fallback)
        _write(fallback)
        return fallback


def build_report(df: pd.DataFrame, long_df: pd.DataFrame) -> Dict[str, pd.DataFrame]:
    n_rows = len(df)
    n_as400 = int(sum(_is_as400_table(norm(t)) for t in df[TABLE_COL]))
    summary = lm.layer_summary(long_df, n_rows, STEP_ORDER)
    for step, label in STEP_LABELS.items():
        summary.loc[summary["layer_key"] == step, "Layer"] = label
    ranking = lm.ranking_table(summary, layer_order=STEP_ORDER).drop(columns=["layer_key"])
    return {
        "Ranking": ranking,
        "Waterfall": lm.waterfall(long_df, n_rows, order=STEP_ORDER),
        "Resolver Breakdown": lm.resolver_breakdown(long_df, n_rows, n_as400, STEP_ORDER),
        "Results": wide_results(df, long_df),
        "Step Detail": step_detail(df, long_df),
        "Notes": pd.DataFrame(NOTES, columns=["Topic", "Description"]),
    }


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def _resolve_input(arg: Optional[str]) -> Path:
    if arg:
        return Path(arg) if Path(arg).is_absolute() else BASE_DIR / arg
    # Checked in order: project root first (back-compat with older layouts),
    # then data/ (the current convention -- test-set xlsx/csv files live
    # there now, not loose in the project root).
    search_dirs = [BASE_DIR, BASE_DIR / "data"]
    for directory in search_dirs:
        for name in DEFAULT_INPUT_CANDIDATES:
            if (directory / name).exists():
                return directory / name
    raise FileNotFoundError(
        f"No input file given and none of the default names exist in {search_dirs} "
        f"({DEFAULT_INPUT_CANDIDATES}). Pass --input \"<your file>.xlsx\" "
        f"(or \"data\\<your file>.xlsx\")."
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Flat 7-step resolver hierarchy evaluation (v3)")
    parser.add_argument("--input", default=None)
    parser.add_argument("--output", default=None)
    parser.add_argument("--sample", type=int, default=None)
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--resume-raw", default=None)
    args = parser.parse_args()

    stamp = f"{dt.datetime.now():%Y%m%d_%H%M%S}"
    out_dir = BASE_DIR / "evaluation"
    out_dir.mkdir(parents=True, exist_ok=True)
    output = Path(args.output) if args.output else out_dir / f"hierarchy_v3_{dt.datetime.now():%Y%m%d}.xlsx"

    provider = os.getenv("LLM_PROVIDER", "bedrock")
    logger.info("LLM_PROVIDER=%s | judge deployment=%s", provider, JUDGE_DEPLOYMENT)

    if args.resume_raw:
        raw = json.loads(Path(args.resume_raw).read_text(encoding="utf-8"))
        df = pd.DataFrame(raw["rows"])
        results = raw["results"]
        logger.info("Resumed %d rows from %s (generation skipped).", len(df), args.resume_raw)
    else:
        input_path = _resolve_input(args.input)
        logger.info("Reading %s ...", input_path)
        df = pd.read_excel(input_path)
        missing = [c for c in (TABLE_COL, COLUMN_COL, GT_TITLE_COL, GT_DESC_COL) if c not in df.columns]
        if missing:
            raise ValueError(f"{input_path.name} is missing columns {missing}. Found: {list(df.columns)}")
        df = df.dropna(subset=[TABLE_COL, COLUMN_COL]).reset_index(drop=True)
        if args.sample and args.sample < len(df):
            df = df.sample(n=args.sample, random_state=42).reset_index(drop=True)
            logger.info("Sampled %d rows.", len(df))
        else:
            logger.info("Running all %d rows.", len(df))

        settings = config.default_settings()
        llm = LLMGeneration(get_llm_client())
        logger.info("Generator client: %s", type(llm.client).__name__)
        deps = Deps(BM25Client(), llm, settings, KataDatasetPool())
        results = run_all_rows(df, deps, args.workers)

        raw_path = out_dir / f"hierarchy_v3_raw_{stamp}.json"
        raw_path.write_text(
            json.dumps({"rows": json.loads(df.to_json(orient="records", force_ascii=False)), "results": results},
                       ensure_ascii=False, default=str),
            encoding="utf-8",
        )
        logger.info("Raw run saved -> %s  (re-judge later with --resume-raw)", raw_path)

    long_df = build_long_df(df, results)
    logger.info("Judging similarity (%s) ...", JUDGE_DEPLOYMENT)
    long_df = judge_similarity(long_df)
    logger.info("Judging truthfulness (kata_similar only, %s) ...", JUDGE_DEPLOYMENT)
    long_df = judge_truthfulness(long_df, df)

    report = build_report(df, long_df)
    logger.info("\n=== RANKING ===\n%s", report["Ranking"].to_string(index=False))
    logger.info("\n=== WATERFALL ===\n%s", report["Waterfall"].to_string(index=False))
    saved = write_workbook(output, report)
    logger.info("Saved -> %s", saved)


if __name__ == "__main__":
    main()